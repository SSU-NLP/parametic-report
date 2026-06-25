#!/usr/bin/env python3
"""지금까지의 모든 transplant 실험을 시트별 엑셀로 정리.

데이터 원본:
  - 단순 spot ablation: reports/transplant_ablation/{cowork,instruct}/*.json  (원본 summary)
  - Bridge 재현:        reports/bridge_repro/s{1024,10000}/*.json
  - k스윕/제외/통제:     reports/results_sequential.csv  (순차 분기점 박제)

usage: python3 scripts/transplant/build_excel.py  (repo 루트에서)
출력: reports/experiments.xlsx
"""
import csv
import json
import os

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

R = "reports"
HDR_FILL = PatternFill("solid", fgColor="DDDDDD")


def add_sheet(wb, name, headers, rows):
    ws = wb.create_sheet(name)
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = HDR_FILL
    for r in rows:
        ws.append(r)
    for i, h in enumerate(headers, 1):
        w = max(len(str(h)), *(len(str(r[i - 1])) for r in rows)) if rows else len(str(h))
        ws.column_dimensions[chr(64 + i) if i <= 26 else "A"].width = min(w + 2, 60)
    ws.freeze_panes = "A2"
    return ws


def jload(path):
    return json.load(open(path)) if os.path.exists(path) else None


def main():
    wb = Workbook()
    wb.remove(wb.active)

    # ── 0. README/맥락 시트 ──
    add_sheet(wb, "_README", ["항목", "내용"], [
        ["실험", "Qwen2.5-1.5B(base) ← Qwen2.5-Coder-1.5B(coder) 파라미터 이식"],
        ["평가", "MultiPL-E HumanEval-Java (completion, 158문제, greedy pass@1)"],
        ["기준", "base=0.2595(41/158), coder=0.3734(59/158, +11.4pp)"],
        ["시트1 ablation_cowork", "단순 spot(|grad·param| top1%) 이식 → 전부 붕괴 (완성=동료 셋업)"],
        ["시트2 ablation_instruct", "단순 spot, instruct+BigCode synthesize (ppl+pass@1)"],
        ["시트3 bridge_repro", "정식 차집합 spot(A∖B) bridge 이식, v1/v2/v3 × sample"],
        ["시트4 bridge_ksweep", "A=B=top-k 키움, v2 교집합 (순차)"],
        ["시트5 bridge_exclusion", "A=1% 고정, 코어 B만 키워 제외 (순차)"],
        ["시트6 bridge_control", "위치/값/인과 통제: rand/perm/ndhi/ndlo/vhi/vlo/reverse (순차)"],
        ["주의", "k스윕/제외/통제는 순차(batch=1). 배치 분기점(2026-06-22) 이후 재측정 예정"],
        ["주의2", "통제 결론: bridge가 random과 pass@1·status 구별 안 됨 (위치 특이성 불지지)"],
        ["", ""],
        ["── strategy/condition 용어 ──", "(이식 = base에 coder 값 덮어쓰기, per-tensor top-k 매칭)"],
        ["base", "원본 base 모델(Qwen2.5-1.5B), 이식 없음 = floor"],
        ["coder", "원본 coder(Qwen2.5-Coder-1.5B) = ceiling(상한)"],
        ["v1", "coder bridge 위치 그대로 이식(같은 위치)"],
        ["v2", "base_bridge ∩ coder_bridge 교집합에 coder 값 (Bridge 핵심)"],
        ["v3", "base bridge 위치에 coder 값"],
        ["v3a/v3b/v3c/v3d", "(ablation) 텐서 내 정렬 대응: 중요도/weight/인덱스순/랜덤짝"],
        ["v3ctrl", "(ablation) base spot 위치 + coder의 non-spot 값"],
        ["rand", "v2와 같은 개수를 무작위 위치에 coder 값 (위치 특이성 통제)"],
        ["perm", "bridge 위치는 그대로, coder 값을 텐서 내 셔플(값-위치 정합 깸)"],
        ["ndhi", "non-bridge 중 |coder−base| drift 큰 곳 N개 (drift confound 통제)"],
        ["ndlo", "non-bridge 중 drift 작은 곳 N개"],
        ["vhi", "bridge 내부에서 drift 상위 절반"],
        ["vlo", "bridge 내부에서 drift 하위 절반"],
        ["reverse", "coder에 base bridge 이식(인과 검증; coder 하락 기대)"],
        ["", ""],
        ["── 컬럼 용어 ──", ""],
        ["k", "A=|weight| top-k 와 B=grad top-k 의 비율 (예: 0.05 = 5%)"],
        ["core_k", "제외 비율 스윕: A는 top1% 고정, 코어 B(grad top-core_k) 크기만 변화"],
        ["selected", "실제 이식된 파라미터 개수"],
        ["decode", "seq=순차(batch=1) / batch=배치(batch=32, deterministic). 둘 사이 ±2문제 디코딩 노이즈"],
        ["pass@1", "HumanEval-Java 158문제 중 통과 비율 (greedy)"],
        ["compile_err / runtime_err", "컴파일 실패 / 실행 후 assert 실패(=wrong answer)"],
    ])

    # ── 1. ablation cowork (생모델 completion) ──
    rows = []
    for c in ["base", "coder", "v1", "v2", "v3a", "v3b", "v3c", "v3d", "v3ctrl"]:
        d = jload(f"{R}/transplant_ablation/cowork/{c}.json")
        if d:
            rows.append([c, d.get("pass_at_1"), d.get("passed"), d.get("total"),
                         d.get("compile_error"), d.get("runtime_error"), d.get("runtime_timeout")])
    add_sheet(wb, "1_ablation_cowork",
              ["condition", "pass@1", "passed", "total", "compile_err", "runtime_err", "timeout"], rows)

    # ── 2. ablation instruct (BigCode synthesize) ──
    rows = []
    for c in ["base-it", "coder-it", "v1-it", "v2-it", "v3a-it", "v3b-it", "v3c-it", "v3d-it", "v3ctrl-it"]:
        d = jload(f"{R}/transplant_ablation/instruct/{c}.json")
        if d:
            rows.append([c, d.get("ppl", {}).get("ppl"), d.get("humaneval", {}).get("pass@1"),
                         d.get("transplant", {}).get("k"), d.get("transplant", {}).get("transplanted_tensors")])
    add_sheet(wb, "2_ablation_instruct", ["condition", "ppl", "pass@1", "k", "transplanted_tensors"], rows)

    # ── 3. bridge repro ──
    rows = []
    for s in ["1024", "10000"]:
        for c in ["base", "coder", "v1", "v2", "v3"]:
            d = jload(f"{R}/bridge_repro/s{s}/{c}.json")
            if d:
                rows.append([s, c, d.get("pass_at_1"), d.get("passed"), d.get("total"),
                             d.get("compile_error"), d.get("runtime_error"), d.get("runtime_timeout")])
    add_sheet(wb, "3_bridge_repro",
              ["sample", "condition", "pass@1", "passed", "total", "compile_err", "runtime_err", "timeout"], rows)

    # ── 4/5/6. CSV에서 ksweep/exclusion/control ──
    csv_rows = list(csv.DictReader(open(f"{R}/results_sequential.csv")))

    def csv_sheet(name, exp):
        rows = []
        for r in csv_rows:
            if r["experiment"] == exp:
                rows.append([r["k"], r["core_k"], r["strategy"], r["decode"], r["pass_at_1"],
                             r["passed"], r["compile_error"], r["runtime_error"], r["timeout"],
                             r["selected"], r["notes"]])
        add_sheet(wb, name, ["k", "core_k", "strategy", "decode", "pass@1", "passed",
                             "compile_err", "runtime_err", "timeout", "selected", "notes"], rows)

    csv_sheet("4_bridge_ksweep", "ksweep")
    csv_sheet("5_bridge_exclusion", "exclusion")
    csv_sheet("6_bridge_control", "control")

    # ── 7/8/9. 배치(분기점 이후) 시트 + origin + 순차vs배치 비교 ──
    bpath = f"{R}/results_batch.csv"
    if os.path.exists(bpath):
        batch_rows = list(csv.DictReader(open(bpath)))

        def bsheet(name, exp):
            rows = [[r["k"], r["core_k"], r["strategy"], r["pass_at_1"], r["passed"],
                     r["compile_error"], r["runtime_error"], r["timeout"], r["selected"], r["notes"]]
                    for r in batch_rows if r["experiment"] == exp]
            add_sheet(wb, name, ["k", "core_k", "strategy", "pass@1", "passed",
                                 "compile_err", "runtime_err", "timeout", "selected", "notes"], rows)

        bsheet("7_batch_ksweep", "ksweep")
        bsheet("8_batch_exclusion", "exclusion")
        bsheet("9_batch_control", "control")
        orows = [[r["strategy"], r["pass_at_1"], r["passed"], r["compile_error"],
                  r["runtime_error"], r["notes"]]
                 for r in batch_rows if r["experiment"] == "origin"]
        add_sheet(wb, "9b_batch_origin",
                  ["strategy", "pass@1", "passed", "compile_err", "runtime_err", "notes"], orows)

        # 10. 순차 vs 배치 (같은 k·strategy의 pass@1 나란히)
        def key(r):
            return (r["experiment"], r["k"], r["core_k"], r["strategy"])
        seqd = {key(r): r["pass_at_1"] for r in csv_rows}
        cmp_rows = []
        for r in batch_rows:
            sk = seqd.get(key(r))
            if sk:
                d = round(float(r["pass_at_1"]) - float(sk), 4)
                cmp_rows.append([r["experiment"], r["k"], r["core_k"], r["strategy"],
                                 sk, r["pass_at_1"], d])
        add_sheet(wb, "10_seq_vs_batch",
                  ["experiment", "k", "core_k", "strategy", "seq_pass@1", "batch_pass@1", "diff"], cmp_rows)

    # ── 11. McNemar (Phase A: base 대비 per-problem flip 검정) ──
    mpath = f"{R}/mcnemar_k0.05.csv"
    if os.path.exists(mpath):
        mrows = list(csv.DictReader(open(mpath)))
        add_sheet(wb, "11_mcnemar", ["ref", "strat", "net(c-b)", "chi2", "p", "verdict"],
                  [[r["ref"], r["strat"], r["net_c_minus_b"], r["chi2"], r["p"], r["verdict"]] for r in mrows])

    # ── 12. Phase B 논문 재현 (3모델 × HumanEval + 일반 5종, code spot damage k-sweep) ──
    try:
        prows = list(csv.reader(open(f"{R}/paper_repro_table1.csv")))
        add_sheet(wb, "12_paper_repro_B", prows[0], prows[1:])
    except FileNotFoundError:
        pass

    # ── 13. Phase C 이식 (qwen base<-coder): Spot/Bridge x 전략 + A-sweep, Python HumanEval chat ──
    # floor(base)=0.5427, ceiling(coder)=0.628. 결론: 전 kA에서 v2<=rand, floor 초과 없음 → bridge 음성.
    try:
        trows = list(csv.reader(open(f"{R}/paper_repro_transplant.csv")))
        add_sheet(wb, "13_transplant_C_py", trows[0], trows[1:])
    except FileNotFoundError:
        pass

    # ── 14. Phase C 이식 — in-distribution Java (humanevalpack) ──
    # 결론: Java v2도 ~floor (ceiling 0.500 미도달). chat(instruction_tokens)도 동일. in-dist 전이도 없음.
    try:
        jrows = list(csv.reader(open(f"{R}/paper_repro_transplant_java.csv")))
        add_sheet(wb, "14_transplant_java", jrows[0], jrows[1:])
    except FileNotFoundError:
        pass

    # ── 15. Python(held-out) vs Java(in-dist) v2 직접 대조 — 둘 다 전이 없음 ──
    try:
        crows = list(csv.reader(open(f"{R}/paper_repro_transplant_compare.csv")))
        add_sheet(wb, "15_py_vs_java", crows[0], crows[1:])
    except FileNotFoundError:
        pass

    # ── 16. Phase C-2 base 쌍 (Qwen2.5-1.5B<-Coder, MultiPL-E java completion): spot vs bridge + McNemar ──
    # 결론: SPOT 이식(v1/v2/v3)=인과 붕괴(2~9/158, p<.001, rand보다 압도적 아래) + reverse도 coder 붕괴.
    #       BRIDGE(여집합)=null(v2 vs rand p=.386). spot은 진짜, bridge는 무효과 — base 쌍에서도 doyun 'up'은 노이즈.
    try:
        brows = list(csv.reader(open(f"{R}/paper_repro_basepair.csv")))
        add_sheet(wb, "16_basepair_C2", brows[0], brows[1:])
    except FileNotFoundError:
        pass

    # ── 17/18. Phase C-2 다언어 MultiPL-E (py/java/cpp/js/go, 798쌍): bridge v2 검정력 결판 ──
    # 결론: 검정력 충분(coder vs base +75 p<.001 검출). 그런데 v2 vs base p=.901, v2 vs rand p=.556 → 무효.
    #       798문제(단일 5배)로도 bridge 'up' 못 잡음 = 잘 검정된 NULL. doyun 'up'=노이즈 확정.
    try:
        m1 = list(csv.reader(open(f"{R}/paper_repro_basepair_multipl.csv")))
        add_sheet(wb, "17_multipl_passk", m1[0], m1[1:])
    except FileNotFoundError:
        pass
    try:
        m2 = list(csv.reader(open(f"{R}/paper_repro_basepair_multipl_mcnemar.csv")))
        add_sheet(wb, "18_multipl_mcnemar_cmh", m2[0], m2[1:])
    except FileNotFoundError:
        pass

    # ── 19. Phase C-2 delta 이식 + global interp (positive control), java ──
    # 결론: global interp(positive control) 실패 — base→coder 선형경로 α≥0.3 붕괴(0.076/0.019) →
    #       두 모델 weight-space 비연결 = sparse/full 모두 weight-copy 불가. spot delta: α↑→붕괴, α↓→floor
    #       (skill carrier 아님=equilibrium-critical). bridge delta: 전 α floor. → 전이 없음 확정.
    try:
        drows = list(csv.reader(open(f"{R}/paper_repro_basepair_delta.csv")))
        add_sheet(wb, "19_basepair_delta", drows[0], drows[1:])
    except FileNotFoundError:
        pass

    # ── 20. 모듈 제한 이식 Round 1: FFN-only global interp gate (vs global-ALL 시트19) ──
    # 결론: gate FAILED — FFN-only도 max α pass@1=0.2975 < base+0.04(0.312), α≥0.3 붕괴(compile_err 폭증).
    #       embed/norm/lm_head/attn 제외해도 붕괴 지속 → 비호환은 FFN 블록 내부 (Case 3, 가장 강한 negative).
    #       raw transplant 종결, alignment(LoRA/activation/TIES)로. Round 2(FFN spot/bridge) 미실행.
    try:
        mrows = list(csv.reader(open(f"{R}/paper_repro_basepair_module.csv")))
        add_sheet(wb, "20_module_ffn_gate", mrows[0], mrows[1:])
    except FileNotFoundError:
        pass

    # ── 21. 이식 호환성 진단 Step1+2 (activation map + 단일모듈 NLL), base←coder ──
    # 기준 nll_base=1.1026 nll_coder=1.1236 (coder≈base NLL → NLL은 전이 신호 약함, 파괴여부만).
    # 결론: 중간 레이어(2~25) 단일 MLP swap은 거의 무해(Δnll +0.02~0.06) → §10 FFN-ALL 붕괴는
    #       누적(28층)+경계레이어(0/1/26/27 비호환: L1 cos0.17·norm1.68, L27 rel_delta3.9) 탓.
    #       raw 완전사망 아님 → Step3(경계 제외 gated delta) 여지, but 전이판정은 pass@1로.
    try:
        arows = list(csv.reader(open(f"{R}/paper_repro_actcompat.csv")))
        add_sheet(wb, "21_act_compat", arows[0], arows[1:])
    except FileNotFoundError:
        pass

    # ── 22. boundary-excluded/middle-only FFN interp (compositional 가설), java pass@1 ──
    # 결론: 중간 24층(2-25) full-swap도 0/158 붕괴(§11 단일층 무해와 반대) → 붕괴는 깊이방향 누적.
    #       boundary 개별(0-1, 26-27)도 붕괴. 중간 α-sweep은 단조 손상, base(0.272) 못 넘음(전이 없음).
    #       교훈: NLL은 generation 붕괴 못 잡음(pass@1 필수). → raw weight transplant 종료, LoRA distill로.
    try:
        brows = list(csv.reader(open(f"{R}/paper_repro_basepair_boundary.csv")))
        add_sheet(wb, "22_boundary_ffn", brows[0], brows[1:])
    except FileNotFoundError:
        pass

    # ── 23. residual-gated donor MLP injection (output-space, L2-25), java pass@1 + McNemar ──
    # 결론: 전 β 유의 향상 없음 — 작은 β(0.01~0.1) neutral(p>0.4, =base), β=0.3 유의 손상(net-12 p=.025).
    #       단 weight-interp(§12 α=0.1→0.247)보다 안정(β=0.1→0.291≈base) → weight 추가손상=MLP 내부 nonlinear path.
    #       donor FFN 출력 방향에 전이 신호 없음 → output-space도 종료 → LoRA distillation으로.
    try:
        rrows = list(csv.reader(open(f"{R}/paper_repro_resgate.csv")))
        add_sheet(wb, "23_resgate", rrows[0], rrows[1:])
    except FileNotFoundError:
        pass

    # ── 24. compatibility-gated static surgery (마지막 무학습 이식), FFN 2-25, java pass@1 + McNemar ──
    # 결론: T(gradient-aligned·Fisher-safe·sign-consistent)·gain·descent-rand·matched-rand 전부 pass@1≈base
    #       (0.272~0.285, 전부 McNemar p≥0.48 무의미). gain은 pred_ΔL=-30.8(코드 loss 큰 감소 예측)인데도 pass@1=base
    #       = 1차 loss-descent ≠ pass@1 전이. → 무학습 parameter transplant 이 모델쌍에서 종결.
    try:
        grows = list(csv.reader(open(f"{R}/paper_repro_gated.csv")))
        add_sheet(wb, "24_gated_surgery", grows[0], grows[1:])
    except FileNotFoundError:
        pass

    # ── 25. abs-accumulated saliency 이식 closure (signed/abs/∩/∖, FFN 2-25) ──
    # 결론: signed≈abs≈signed∩abs≈abs∖signed≈matched-rand≈base (0.26~0.285, McNemar p≥0.48 전부 무의미).
    #       saliency 정의(signed vs abs activity)를 바꿔도 static transplant 차이 0 → static saliency surgery 완전 종결.
    try:
        srows = list(csv.reader(open(f"{R}/paper_repro_saliency.csv")))
        add_sheet(wb, "25_saliency", srows[0], srows[1:])
    except FileNotFoundError:
        pass

    out = f"{R}/experiments.xlsx"
    wb.save(out)
    print(f"saved {out} ({len(wb.sheetnames)} sheets: {wb.sheetnames})")


if __name__ == "__main__":
    main()
